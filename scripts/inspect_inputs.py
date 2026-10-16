#!/usr/bin/env python3
"""Extract local PDF evidence and an optional DOCX formatting profile."""
import argparse
import json
from pathlib import Path
import re


def template_profile(path):
    from docx import Document
    from lxml import etree
    doc = Document(str(path))
    def xml(node):
        return None if node is None else etree.tostring(node, encoding="unicode")
    sections = []
    for s in doc.sections:
        sections.append({"page_width": s.page_width, "page_height": s.page_height,
                         "top_margin": s.top_margin, "bottom_margin": s.bottom_margin,
                         "left_margin": s.left_margin, "right_margin": s.right_margin,
                         "header_distance": s.header_distance, "footer_distance": s.footer_distance,
                         "different_first_page": s.different_first_page_header_footer,
                         "xml": xml(s._sectPr)})
    return {"source": path.name, "units": "Lengths are EMU; raw OOXML attributes retain their native units.",
            "caveat": "Direct properties may be unset. Effective appearance also depends on style inheritance, theme, document defaults, fonts, and renderer. Inspect rendered reference pages.",
            "sections": sections,
            "styles": [{"id": s.style_id, "name": s.name, "base": s.base_style.style_id if s.base_style else None,
                        "xml": xml(s.element)} for s in doc.styles],
            "paragraphs": [{"index": i, "text": p.text, "style": p.style.style_id,
                            "pPr": xml(p._p.pPr), "runs": [{"text": r.text, "rPr": xml(r._r.rPr)} for r in p.runs]}
                           for i, p in enumerate(doc.paragraphs)],
            "tables": [{"index": i, "rows": len(t.rows), "columns": len(t.columns),
                        "text": [[c.text for c in row.cells] for row in t.rows], "xml": xml(t._tbl)}
                       for i, t in enumerate(doc.tables)]}


def render_source(path, out):
    try:
        import pymupdf
        doc = pymupdf.open(str(path))
        for i, page in enumerate(doc):
            page.get_pixmap(matrix=pymupdf.Matrix(0.8, 0.8), alpha=False).save(str(out / f"source-{i+1:03d}.png"))
        doc.close()
        return "PyMuPDF"
    except ImportError:
        import pypdfium2 as pdfium
        doc = pdfium.PdfDocument(str(path))
        for i in range(len(doc)):
            page = doc[i]
            bitmap = page.render(scale=0.8)
            bitmap.to_pil().save(out / f"source-{i+1:03d}.png")
            bitmap.close()
            page.close()
        doc.close()
        return "PDFium (fallback; compare text glyphs with source)"


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pdf", type=Path)
    parser.add_argument("--out", required=True, type=Path)
    parser.add_argument("--template", type=Path)
    parser.add_argument("--render", action="store_true")
    args = parser.parse_args()
    args.out.mkdir(parents=True, exist_ok=True)
    import pdfplumber
    pages, candidates, suspects = [], [], []
    with pdfplumber.open(args.pdf) as pdf:
        for i, page in enumerate(pdf.pages, 1):
            text = page.extract_text() or ""
            lines = page.extract_text_lines(layout=False, return_chars=True)
            pages.append({"page": i, "width": page.width, "height": page.height,
                          "text": text, "lines": lines, "chars": page.chars})
            if len(text.strip()) < 30:
                suspects.append({"page": i, "reason": "Little extractable text: inspect for scanned page or image-only equations; OCR is not automatic."})
            for line in lines:
                value = line.get("text", "")
                if re.search(r"[∑∫∏√=]|[（(]\s*\d+\.\d+\s*[）)]\s*$", value):
                    candidates.append({"page": i, "text": value,
                                       "bbox": [line.get(k) for k in ("x0", "top", "x1", "bottom")]})
    data = {"source": args.pdf.name, "page_count": len(pages), "pages": pages,
            "formula_candidates": candidates, "ocr_suspects": suspects,
            "warning": "Formula candidates are heuristic evidence, not a complete formula inventory or trustworthy transcription."}
    (args.out / "pdf-evidence.json").write_text(json.dumps(data, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    if args.template:
        (args.out / "template-profile.json").write_text(json.dumps(template_profile(args.template), ensure_ascii=False, indent=2), encoding="utf-8")
    renderer = render_source(args.pdf, args.out) if args.render else None
    print(json.dumps({"pages": len(pages), "formula_candidates": len(candidates), "ocr_suspects": suspects,
                      "renderer": renderer, "output": str(args.out.resolve())}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
