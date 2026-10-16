"""Build editable Word from an agent-verified reconstruction plan, not raw PDF text."""
import argparse
from copy import deepcopy
import json
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.text.paragraph import Paragraph
from native_math import display_equation, latex_to_omml


def node(tag, **attrs):
    e = OxmlElement('w:' + tag)
    for k, v in attrs.items():
        e.set(qn('w:' + k), str(v))
    return e


def font(run, spec):
    if 'latin' in spec:
        run.font.name = spec['latin']
    if 'font' in spec:
        run._r.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), spec['font'])
    for key in ('bold', 'italic', 'subscript', 'superscript'):
        if key in spec:
            setattr(run.font, key, spec[key])
    if 'size' in spec:
        run.font.size = Pt(spec['size'])
    if 'color' in spec:
        run.font.color.rgb = RGBColor.from_string(spec['color'].lstrip('#'))


def para_format(p, spec):
    if 'align' in spec:
        p.alignment = {'left': 0, 'center': 1, 'right': 2, 'justify': 3}[spec['align']]
    for key in ('space_before', 'space_after', 'first_line_indent'):
        if key in spec:
            setattr(p.paragraph_format, key, Pt(spec[key]))
    if 'line_spacing' in spec:
        p.paragraph_format.line_spacing = spec['line_spacing']
    if 'line_exact' in spec:
        p.paragraph_format.line_spacing = Pt(spec['line_exact'])
    if 'new_page' in spec:
        p.paragraph_format.page_break_before = spec['new_page']


def add_field(p, instruction, cached='?'):
    f = node('fldSimple', instr=instruction)
    r = node('r')
    t = node('t'); t.text = cached; r.append(t); f.append(r)
    p._p.append(f)
    return f


class Builder:
    def __init__(self, plan, template=None, base_dir=None):
        self.plan = plan
        self.base_dir = Path(base_dir or '.')
        self.doc = Document(template) if template else Document()
        self.patterns = {}
        paragraphs = self.doc.paragraphs
        for role, index in plan.get('template_roles', {}).items():
            if not isinstance(index, int) or not 0 <= index < len(paragraphs):
                raise ValueError(f'Invalid template paragraph for {role}: {index}')
            self.patterns[role] = deepcopy(paragraphs[index]._p)
        self.formats = plan.get('format', {})
        self.template = bool(template)
        self.chapter_style_name = 'Fidelity Chapter'
        if 'chapter' not in self.patterns and self.chapter_style_name not in self.doc.styles:
            style = self.doc.styles.add_style(self.chapter_style_name, WD_STYLE_TYPE.PARAGRAPH)
            style.base_style = self.doc.styles['Heading 1']
        body = self.doc.element.body
        section = deepcopy(body.sectPr)
        for child in list(body):
            body.remove(child)
        if section is not None:
            body.append(section)
        self.body = body
        self.heading_entries = []
        for i, block in enumerate(plan.get('blocks', [])):
            if block['type'] == 'heading':
                self.heading_entries.append((block, f'_Fidelity_{i}'))
        self.anchors = {id(b): a for b, a in self.heading_entries}
        page = plan.get('page', {})
        for s in self.doc.sections:
            for key, value in page.items():
                attr = {'width': 'page_width', 'height': 'page_height'}.get(key, key + '_margin')
                if attr not in ('page_width', 'page_height', 'top_margin', 'bottom_margin', 'left_margin', 'right_margin'):
                    raise ValueError(f'Unknown page property: {key}')
                setattr(s, attr, Pt(value))

    def paragraph(self, role='body', content=None, **overrides):
        pattern = self.patterns.get(role)
        if pattern is not None:
            xml = deepcopy(pattern)
            for c in list(xml):
                if c.tag != qn('w:pPr'):
                    xml.remove(c)
            pp = xml.find(qn('w:pPr'))
            if pp is not None:
                for name in ('sectPr', 'numPr'):
                    for c in pp.findall(qn('w:' + name)):
                        pp.remove(c)
            self.body.insert(len(self.body) - 1, xml)
            p = Paragraph(xml, self.doc._body)
        else:
            p = self.doc.add_paragraph()
            if role in ('chapter', 'h2', 'h3'):
                p.style = {'chapter': self.chapter_style_name, 'h2': 'Heading 2', 'h3': 'Heading 3'}[role]
        spec = dict(self.formats.get(role, {})); spec.update(overrides)
        if not self.template:
            defaults = {'font': '宋体', 'latin': 'Times New Roman', 'size': 12, 'color': '000000'}
            defaults.update(spec); spec = defaults
        para_format(p, spec)
        if content is not None:
            self.runs(p, content, role, spec)
        return p

    def runs(self, p, content, role='body', spec=None):
        items = [{'text': content}] if isinstance(content, str) else content
        for item in items:
            if 'math' in item:
                p._p.append(latex_to_omml(item['math']))
                continue
            r = p.add_run(item.get('text', ''))
            pattern = self.patterns.get(role)
            if pattern is not None:
                rp = pattern.find('.//' + qn('w:rPr'))
                if rp is not None:
                    old = r._r.find(qn('w:rPr'))
                    if old is not None:
                        r._r.remove(old)
                    r._r.insert(0, deepcopy(rp))
            font(r, dict(spec or self.formats.get(role, {}), **{k: v for k, v in item.items() if k != 'text'}))

    def heading(self, block):
        role = block.get('role', {1: 'chapter', 2: 'h2', 3: 'h3'}.get(block.get('level', 1)))
        if role not in ('chapter', 'h2', 'h3'):
            raise ValueError('Heading role must be chapter, h2, or h3')
        p = self.paragraph(role, block['text'], **block.get('format', {}))
        outline = p._p.get_or_add_pPr().find(qn('w:outlineLvl'))
        if outline is None:
            outline = node('outlineLvl'); p._p.pPr.append(outline)
        outline.set(qn('w:val'), str({'chapter': 0, 'h2': 1, 'h3': 2}[role]))
        anchor = self.anchors[id(block)]
        index = str(len(self.doc.element.xpath('//w:bookmarkStart')) + 1000)
        p._p.insert(1 if p._p.pPr is not None else 0, node('bookmarkStart', id=index, name=anchor))
        p._p.append(node('bookmarkEnd', id=index))

    def table(self, block):
        rows = block['rows']
        if not rows or not rows[0] or any(len(r) != len(rows[0]) for r in rows):
            raise ValueError('Table rows must be nonempty and rectangular')
        if block.get('caption'):
            p = self.paragraph('caption', block['caption'])
            p.paragraph_format.keep_with_next = True
        t = self.doc.add_table(rows=0, cols=len(rows[0])); t.autofit = False
        t._tbl.set(qn('w:rsidR'), '00000000')
        widths = block.get('widths')
        if widths is None:
            sec = self.doc.sections[0]
            width = (sec.page_width - sec.left_margin - sec.right_margin) / 12700
            widths = [width / len(rows[0])] * len(rows[0])
        if len(widths) != len(rows[0]) or any(w <= 0 for w in widths):
            raise ValueError('Widths must match the number of columns')
        for col, w in zip(t.columns, widths):
            col.width = Pt(w)
        pr = t._tbl.tblPr
        for name in ('tblBorders', 'tblStyle', 'tblLook'):
            for old in pr.findall(qn('w:' + name)):
                pr.remove(old)
        borders = node('tblBorders')
        for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            borders.append(node(side, val='nil' if block.get('three_line', False) else 'single', sz=4, color='000000'))
        pr.append(borders)
        for ri, row in enumerate(rows):
            rr = t.add_row()
            if ri == 0:
                rr._tr.get_or_add_trPr().append(node('tblHeader'))
            rr._tr.get_or_add_trPr().append(node('cantSplit'))
            for ci, value in enumerate(row):
                cell = rr.cells[ci]; cell.width = Pt(widths[ci])
                if block.get('three_line', False):
                    cb = node('tcBorders')
                    for side in ('top', 'bottom', 'left', 'right'):
                        visible = (ri == 0 and side in ('top', 'bottom')) or (ri == len(rows) - 1 and side == 'bottom')
                        cb.append(node(side, val='single' if visible else 'nil', sz=8 if ri == 0 and side == 'top' else 4, color='000000'))
                    cell._tc.get_or_add_tcPr().append(cb)
                p = cell.paragraphs[0]; p.alignment = 0
                p.paragraph_format.first_line_indent = Pt(0)
                p.paragraph_format.space_after = Pt(2)
                self.runs(p, str(value), spec={'font': block.get('font', '宋体'), 'latin': 'Times New Roman', 'size': block.get('size', 10.5), 'bold': ri == 0})
        return t

    def contents(self):
        spec = self.plan.get('toc')
        if not spec:
            return
        p = self.paragraph('toc_title', spec.get('title', '目录'), font='宋体', latin='Times New Roman', size=12, align='left', space_after=14)
        p._p.get_or_add_pPr().append(node('outlineLvl', val=9))
        if spec.get('layout', 'single') == 'single':
            p = self.paragraph('toc')
            add_field(p, ' TOC \\o "1-3" \\h \\z \\u ', 'Update fields in Word')
            return
        sec = self.doc.sections[0]
        width = (sec.page_width - sec.left_margin - sec.right_margin) / 12700
        t = self.doc.add_table(rows=1, cols=2); t.autofit = False
        pr = t._tbl.tblPr; pr.append(node('tblCaption', val='FidelityTOC'))
        borders = node('tblBorders')
        for side in ('top', 'bottom', 'left', 'right', 'insideH', 'insideV'):
            borders.append(node(side, val='nil'))
        pr.append(borders)
        split = spec.get('split_after', (len(self.heading_entries) + 1) // 2)
        for ci, entries in enumerate((self.heading_entries[:split], self.heading_entries[split:])):
            cell = t.cell(0, ci); cell.width = Pt(width / 2)
            cp = cell._tc.get_or_add_tcPr(); cp.append(node('vAlign', val='top'))
            margins = node('tcMar'); cp.append(margins)
            gap = spec.get('gap', 24)
            margins.append(node('left', w=gap * 10 if ci else 0, type='dxa'))
            margins.append(node('right', w=0 if ci else gap * 10, type='dxa'))
            for j, (block, anchor) in enumerate(entries):
                p = cell.paragraphs[0] if j == 0 else cell.add_paragraph()
                depth = {'chapter': 0, 'h2': 1, 'h3': 2}.get(block.get('role'), max(0, block.get('level', 1) - 1))
                pp = p._p.get_or_add_pPr()
                pp.append(node('outlineLvl', val=9)); pp.append(node('ind', left=depth * 160, right=240, firstLine=0))
                tabs = node('tabs'); tabs.append(node('tab', val='right', leader='dot', pos=int((width / 2 - gap / 2) * 20))); pp.append(tabs)
                p.paragraph_format.line_spacing = Pt(spec.get('line', 16.5)); p.paragraph_format.space_after = Pt(3)
                h = node('hyperlink', anchor=anchor, history=1); p._p.append(h)
                r = node('r'); h.append(r)
                rp = node('rPr'); r.append(rp)
                rp.append(node('rFonts', ascii='Times New Roman', hAnsi='Times New Roman', eastAsia='宋体'))
                rp.append(node('sz', val=int(spec.get('size', 10.5) * 2))); rp.append(node('color', val='000000'))
                rp.append(node('b', val=1 if depth == 0 else 0))
                text = node('t'); text.text = block['text']; r.append(text)
                tab = node('r'); tab.append(node('tab')); h.append(tab)
                f = node('fldSimple', instr=f' PAGEREF {anchor} \\h ')
                rr = node('r'); rr.append(deepcopy(rp)); tt = node('t'); tt.text = '?'; rr.append(tt); f.append(rr); h.append(f)

    def headers(self):
        spec = self.plan.get('headers')
        if not spec:
            return
        self.doc.settings.odd_and_even_pages_header_footer = True
        for s in self.doc.sections:
            s.different_first_page_header_footer = spec.get('first_page_blank', False)
            for header, value, align in ((s.header, spec.get('odd', ''), 2), (s.even_page_header, spec.get('even', ''), 0)):
                for child in list(header._element):
                    header._element.remove(child)
                p = header.add_paragraph(); p.alignment = align
                p.paragraph_format.first_line_indent = Pt(0)
                if value == 'chapter':
                    chapter = self.patterns.get('chapter')
                    sid = chapter.find('./' + qn('w:pPr') + '/' + qn('w:pStyle')) if chapter is not None else None
                    style = sid.get(qn('w:val')) if sid is not None else self.chapter_style_name
                    style_name = self.doc.styles[style].name if style in self.doc.styles else 'Heading 1'
                    f = add_field(p, f' STYLEREF "{style_name}" ', 'Chapter')
                    r = f.find(qn('w:r'))
                    from docx.text.run import Run
                    font(Run(r, p), {'font': '宋体', 'latin': 'Times New Roman', 'size': 9, 'color': '000000', 'bold': False} | spec.get('format', {}))
                else:
                    r = p.add_run(value); font(r, {'font': '宋体', 'latin': 'Times New Roman', 'size': 9, 'color': '000000', 'bold': False} | spec.get('format', {}))

    def build(self, output):
        self.contents()
        for block in self.plan['blocks']:
            kind = block['type']
            if kind == 'heading':
                self.heading(block)
            elif kind == 'paragraph':
                self.paragraph('body', block.get('runs', block.get('text', '')), **block.get('format', {}))
            elif kind == 'equation':
                p = self.paragraph('equation', align='center', first_line_indent=0, space_before=3, space_after=3)
                p._p.append(display_equation(block['latex'], block.get('number'), force_below=block.get('force_below', False)))
            elif kind == 'table':
                self.table(block)
            elif kind == 'image':
                path = (self.base_dir / block['path']).resolve()
                p = self.paragraph('figure', align='center', first_line_indent=0)
                p.add_run().add_picture(str(path), width=Pt(block.get('width', 300)))
                if block.get('caption'):
                    self.paragraph('caption', block['caption'])
            elif kind == 'page_break':
                self.doc.add_page_break()
            else:
                raise ValueError(f'Unsupported block type: {kind}')
        self.headers()
        output = Path(output)
        if output.exists():
            raise FileExistsError(f'Refusing to overwrite: {output}')
        output.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(output)
        return output


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--plan', type=Path, required=True)
    parser.add_argument('--template', type=Path)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    plan = json.loads(args.plan.read_text(encoding='utf-8-sig'))
    print(Builder(plan, args.template, args.plan.parent).build(args.output))
    if plan.get('toc') or plan.get('headers'):
        print('Update fields and render before delivery. This build is not a QA approval.')


if __name__ == '__main__':
    main()
