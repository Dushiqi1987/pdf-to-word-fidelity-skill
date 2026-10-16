import importlib.util
import json
from pathlib import Path
import sys
import tempfile
import unittest
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
from build_docx import Builder


class BuilderTests(unittest.TestCase):
    def test_public_plan_builds_native_editable_artifact(self):
        plan = json.loads((ROOT / 'examples/demo-plan.json').read_text(encoding='utf8'))
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / 'result.docx'
            Builder(plan).build(out)
            d = Document(out)
            self.assertEqual(len(d.element.xpath('//m:oMath')), 4)
            self.assertEqual(len(d.tables), 2)
            self.assertEqual(len(d.inline_shapes), 0)
            self.assertEqual(len(d.tables[0]._tbl.xpath('.//w:fldSimple')), 5)
            self.assertIn('6.16', ''.join(d.element.xpath('//m:t/text()')))
            borders = d.tables[1]._tbl.xpath('.//w:tcBorders/w:left')
            self.assertTrue(borders)
            self.assertTrue(all(b.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val') == 'nil' for b in borders))
            self.assertTrue(d.settings.odd_and_even_pages_header_footer)
            with self.assertRaises(FileExistsError):
                Builder(plan).build(out)

    def test_template_page_and_body_style_are_cloned(self):
        with tempfile.TemporaryDirectory() as tmp:
            template = Path(tmp) / 'template.docx'
            d = Document(); p = d.add_paragraph('Reference')
            p.paragraph_format.space_before = __import__('docx').shared.Pt(11)
            d.sections[0].left_margin = __import__('docx').shared.Pt(83)
            d.save(template)
            out = Path(tmp) / 'new.docx'
            Builder({'template_roles': {'body': 0}, 'blocks': [{'type': 'paragraph', 'text': 'Fresh content'}]}, template).build(out)
            n = Document(out)
            self.assertEqual(n.sections[0].left_margin.pt, 83)
            self.assertEqual(n.paragraphs[0].paragraph_format.space_before.pt, 11)
            self.assertEqual(n.paragraphs[0].text, 'Fresh content')

    def test_invalid_table_is_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            plan = {'blocks': [{'type': 'table', 'rows': [['A', 'B'], ['C']]}]}
            with self.assertRaises(ValueError):
                Builder(plan).build(Path(tmp) / 'bad.docx')


if __name__ == '__main__':
    unittest.main()
